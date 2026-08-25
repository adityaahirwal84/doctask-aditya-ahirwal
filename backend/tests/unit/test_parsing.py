import io

import docx
import fitz

from app.services.parsing import parse_docx, parse_markdown, parse_pdf, parse_txt


def test_parse_txt_splits_on_blank_lines():
    data = b"First paragraph.\n\nSecond paragraph.\n\nThird paragraph."
    chunks = parse_txt(data)
    assert [c.content for c in chunks] == ["First paragraph.", "Second paragraph.", "Third paragraph."]


def test_parse_markdown_tracks_headings_as_section_ref():
    data = b"# Section One\n\nSome content here.\n\n## Subsection\n\nMore content."
    chunks = parse_markdown(data)
    assert chunks[0].section_ref == "Section One"
    assert any(c.section_ref == "Subsection" for c in chunks)


def test_parse_docx_extracts_paragraphs_and_heading_section_refs():
    doc = docx.Document()
    doc.add_heading("Payment Terms", level=1)
    doc.add_paragraph("Payment is due net 30 days from invoice.")
    buf = io.BytesIO()
    doc.save(buf)

    chunks = parse_docx(buf.getvalue())
    contents = [c.content for c in chunks]
    assert "Payment is due net 30 days from invoice." in contents
    matching = [c for c in chunks if c.content == "Payment is due net 30 days from invoice."][0]
    assert matching.section_ref == "Payment Terms"


def test_parse_docx_extracts_table_rows():
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Invoice Number"
    table.rows[0].cells[1].text = "INV-1042"
    buf = io.BytesIO()
    doc.save(buf)

    chunks = parse_docx(buf.getvalue())
    assert any("INV-1042" in c.content for c in chunks)


def test_parse_pdf_extracts_text_with_page_refs():
    pdf_doc = fitz.open()
    page = pdf_doc.new_page()
    page.insert_text((72, 72), "Payment terms are net 30 days from invoice.")
    data = pdf_doc.tobytes()
    pdf_doc.close()

    chunks = parse_pdf(data)
    assert any("Payment terms" in c.content for c in chunks)
    assert chunks[0].section_ref == "page 1"


def test_parse_pdf_separates_visually_distinct_blocks_without_blank_line_markers():
    """Regression test: a PDF built from several directly-positioned text
    runs (common for simple/manual PDF generation, and not unheard of in
    real-world PDFs) has no blank-line paragraph markers in its plain-text
    extraction - flat 'text' mode collapses everything onto single
    newlines. parse_pdf must still separate these into distinct chunks
    using PyMuPDF's genuine visual block detection, not silently return
    the whole page as one oversized chunk."""
    pdf_doc = fitz.open()
    page = pdf_doc.new_page()
    y = 72
    for line in ["INVOICE", "Invoice Number: INV-9001", "Amount Due: $500.00"]:
        page.insert_text((72, y), line, fontsize=11)
        y += 30  # each insert_text call is its own visually distinct block
    data = pdf_doc.tobytes()
    pdf_doc.close()

    chunks = parse_pdf(data)
    assert len(chunks) == 3, f"expected 3 distinct chunks, got {len(chunks)}: {[c.content for c in chunks]}"
    assert chunks[0].content == "INVOICE"
    assert chunks[1].content == "Invoice Number: INV-9001"
    assert chunks[2].content == "Amount Due: $500.00"


def test_long_block_is_split_under_max_chars():
    data = ("word " * 500).encode()  # ~2500 chars, single paragraph
    chunks = parse_txt(data)
    assert len(chunks) > 1
    assert all(len(c.content) <= 1500 for c in chunks)
